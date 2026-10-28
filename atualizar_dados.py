"""Atualiza os dados do app a partir de fontes oficiais do Ministério da Previdência.

Fontes:
  - ISP-RPPS (planilha "resultado final" no gov.br): porte, índice ISP, perfil atuarial, nível Pró-Gestão
    e a posição no ranking (calculada aqui a partir da classificação do ISP).
  - API pública do CADPREV (apicadprev.trabalho.gov.br): DAIR (saldo da carteira) e CRP.

Uso no projeto (gera app/dados/rpps.json):
    pip install openpyxl
    python scripts/atualizar_dados.py

Uso no repositório publicado (troca os dados dentro do index.html, roda todo dia no GitHub Actions):
    python atualizar_dados.py --html index.html

Só regrava quando algo mudou. Se a API do CADPREV falhar, mantém os valores anteriores e avisa.
Também vigia a página do Manual de Certificação: se sair versão maior que a de regras.js, avisa e falha o job.
"""
import json
import re
import sys
import tempfile
import time
import unicodedata
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from pathlib import Path

import openpyxl

PAGINA_ISP = "https://www.gov.br/previdencia/pt-br/assuntos/rpps/indice-de-situacao-previdenciaria"
PAGINA_MANUAL = "https://www.gov.br/previdencia/pt-br/assuntos/rpps/Requisitos-para-Dirigentes-e-Conselheiros-de-RPPS/manual-da-certificacao-profissional"
API = "https://apicadprev.trabalho.gov.br"
MANUAL_NO_APP = (2, 0)  # versão do Manual usada em regras.js (REGRAS.manualVersao)
BASE = "https://www.gov.br"
DESTINO = Path(__file__).resolve().parent.parent / "app" / "dados" / "rpps.json"
UA = {"User-Agent": "Mozilla/5.0 (LEMA Edu - atualizador de dados RPPS)"}
BLOCO = re.compile(r'(<script id="dados" type="application/json">)(.*?)(</script>)', re.S)
FECHA = "<" + "\\" + "/"  # "<\/", forma segura de escrever "</" dentro de <script>
NOTA = {"A": 1.5, "B": 1.0, "C": 0.5, "D": 0.0}  # notas da metodologia do ISP (aba "Combinação de Resultados")


def baixar(url, destino=None, tentativas=3):
    ultimo = None
    for i in range(tentativas):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=180) as r:
                dados = r.read()
            if destino:
                Path(destino).write_bytes(dados)
            return dados
        except Exception as e:  # noqa: BLE001
            ultimo = e
            time.sleep(2 * (i + 1))
    raise ultimo


def api(caminho):
    return json.loads(baixar(API + caminho).decode("utf-8", "replace"))["data"]


def sem_acento(s):
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().upper().strip()


def cnpj14(v):
    return re.sub(r"\D", "", str(v)).zfill(14)


def achar_planilhas_finais():
    """Devolve {ano: url} das planilhas de resultado final publicadas na página do ISP."""
    html = baixar(PAGINA_ISP).decode("utf-8", "ignore")
    links = re.findall(r'href="([^"]+\.xlsx)"', html, flags=re.I)
    achados = {}
    for link in links:
        nome = link.lower()
        if "preliminar" in nome or "previo" in nome:
            continue
        if "final" not in nome and "consolidado" not in nome:
            continue
        ano = re.search(r"(20\d\d)", nome)
        if not ano:
            continue
        url = link if link.startswith("http") else BASE + link
        achados.setdefault(int(ano.group(1)), url)
    return achados


def ler_planilha(caminho):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    linhas = list(ws.iter_rows(values_only=True))
    ini = next(i for i, r in enumerate(linhas) if r and sem_acento(r[0]) == "ENTE")
    cab = [sem_acento(c or "") for c in linhas[ini]]

    def col(*partes):
        for i, c in enumerate(cab):
            if all(p in c for p in partes):
                return i
        raise KeyError(partes)

    def col_opc(*partes):
        try:
            return col(*partes)
        except KeyError:
            return None

    i_ente, i_uf = col("ENTE"), col("UF")
    i_cnpj = col_opc("CNPJ")
    i_grupo, i_sub = col("GRUPO"), col("SUBGRUPO")
    i_isp = col_opc("INDICADOR DE SITUACAO PREVIDENCIARIA")  # a planilha do ano anterior pode ter outro formato
    i_perfil = col("PERFIL ATUARIAL")
    # classificações por dimensão e indicadores individuais (para o ranking)
    i_cls = [col_opc("CLASSIFICACAO EM GESTAO"), col_opc("CLASSIFICACAO EM FINANCAS"), col_opc("CLASSIFICACAO EM ATUARIA")]
    i_ind = [col_opc("INDICADOR DE REGULARIDADE"), col_opc("INDICADOR ENVIO"), col_opc("INDICADOR DE GESTAO"),
             col_opc("SUFICIENCIA"), col_opc("ACUMULACAO"), col_opc("EQUACIONAMENTO"),
             col_opc("COBERTURA"), col_opc("SUSTENTABILIDADE DAS PROVISOES EM RELACAO"), col_opc("REFORMA")]

    def nota(r, i):
        return NOTA.get(str(r[i]).strip()[:1].upper(), 0.0) if i is not None and r[i] is not None else 0.0

    out = {}
    for r in linhas[ini + 1:]:
        if not r[i_ente] or not r[i_uf]:
            continue
        nome = re.sub(r"\s+-\s+[A-Z]{2}$", "", str(r[i_ente]).strip())
        out[(r[i_uf], nome)] = {
            "cnpj": cnpj14(r[i_cnpj]) if i_cnpj is not None and r[i_cnpj] else "",
            "grupo": sem_acento(r[i_grupo] or ""),
            "subgrupo": sem_acento(r[i_sub] or ""),
            "isp": str(r[i_isp]).strip() if i_isp is not None and r[i_isp] else "",
            "perfil": str(r[i_perfil]).strip() if r[i_perfil] else "",
            "score": (sum(nota(r, i) for i in i_cls), sum(nota(r, i) for i in i_ind)),
        }

    # nível do Pró-Gestão (aba GESTÃO)
    pg = {}
    for aba in wb.worksheets:
        if sem_acento(aba.title) == "GESTAO":
            rows = list(aba.iter_rows(values_only=True))
            h = next((i for i, r in enumerate(rows) if r and sem_acento(r[0]) == "ENTE"), None)
            if h is not None:
                hc = [sem_acento(c or "") for c in rows[h]]
                ni = next((i for i, c in enumerate(hc) if "NIVEL DA CERTIFICACAO" in c), None)
                if ni is not None:
                    for r in rows[h + 1:]:
                        if r[0] and r[1]:
                            nome = re.sub(r"\s+-\s+[A-Z]{2}$", "", str(r[0]).strip())
                            v = str(r[ni]).strip() if r[ni] is not None else "-"
                            pg[(r[1], nome)] = v if v in ("I", "II", "III", "IV") else "-"
    for k, v in out.items():
        v["pg"] = pg.get(k, "-")
    return out


def calcular_ranking(atual):
    """Posição no Brasil: ordena pela soma das notas das 3 dimensões do ISP e desempata pela soma dos
    9 indicadores. Empates dividem a mesma posição (a melhor). Cálculo próprio da LEMA Edu."""
    chaves = {k: v["score"] for k, v in atual.items() if v["isp"]}
    ordenadas = sorted(chaves.values(), reverse=True)
    pos = {}
    for k, s in chaves.items():
        pos[k] = 1 + sum(1 for o in ordenadas if o > s)
    return pos, len(chaves)


def versao_do_manual():
    """Maior 'versão X.Y' listada na página oficial do Manual de Certificação."""
    try:
        html = baixar(PAGINA_MANUAL).decode("utf-8", "ignore")
    except Exception as e:  # noqa: BLE001  página fora do ar não deve derrubar a atualização dos dados
        print("::warning::Não consegui abrir a página do Manual:", e)
        return None
    achadas = [(int(a), int(b)) for a, b in re.findall(r"vers[ãa]o\s*(\d+)\.(\d+)", html, flags=re.I)]
    return max(achadas) if achadas else None


# ---------------------------------------------------------------- CADPREV: CRP e DAIR
def iso(dmy):
    try:
        return datetime.strptime(dmy, "%d/%m/%Y").date().isoformat()
    except Exception:  # noqa: BLE001
        return ""


def crp_de(linhas):
    """Escolhe o CRP mais recente de uma lista de linhas da API e devolve um dicionário compacto."""
    if not linhas:
        return None
    melhor = max(linhas, key=lambda x: (iso(x.get("dt_emissao", "")), iso(x.get("dt_validade", ""))))
    validade = iso(melhor.get("dt_validade", ""))
    return {
        "t": "J" if str(melhor.get("ds_situacao", "")).upper().startswith("J") else "A",
        "v": validade, "e": iso(melhor.get("dt_emissao", "")), "n": melhor.get("nr_crp", ""),
        "ok": bool(validade) and validade >= date.today().isoformat(),
    }


def buscar_crp(cnpjs):
    """CRP de cada CNPJ: primeiro em lote os vigentes, depois um a um os que não têm vigente."""
    por_cnpj = {}
    for x in api("/RPPS_CRP?tp_crp=" + urllib.parse.quote("VÁLIDO") + "&limit=200000"):
        por_cnpj.setdefault(cnpj14(x["nr_cnpj_entidade"]), []).append(x)
    res = {c: crp_de(l) for c, l in por_cnpj.items()}
    faltam = [c for c in cnpjs if not res.get(c) or not res[c]["ok"]]

    def um(c):
        try:
            return c, crp_de(api("/RPPS_CRP?nr_cnpj_entidade=" + c))
        except Exception:  # noqa: BLE001
            return c, res.get(c)
    with ThreadPoolExecutor(8) as ex:
        for c, v in ex.map(um, faltam):
            if v:
                res[c] = v
    return res


def ultimo_mes_dair(ano_atual):
    """Último (ano, mês) de DAIR enviado por CNPJ."""
    ult = {}
    for ano in (ano_atual, ano_atual - 1):
        for x in api(f"/DAIR_IDENTIFICACAO?dt_ano={ano}&limit=200000"):
            if not str(x.get("te_finalidade", "")).upper().startswith("ENCERRAMENTO"):
                continue
            c = cnpj14(x["nr_cnpj_entidade"])
            m = (int(x["dt_ano"]), int(x["dt_mes"]))
            if m > ult.get(c, (0, 0)):
                ult[c] = m
    return ult


def saldo_dair(cnpj, ano, mes):
    """Soma da carteira do DAIR sem 'Disponibilidades Financeiras' e sem 'Imóveis' (confere com o prorpps)."""
    linhas = api(f"/DAIR_CARTEIRA?nr_cnpj_entidade={cnpj}&dt_ano={ano}&dt_mes_bimestre={mes}&limit=200000")
    total = 0.0
    for x in linhas:
        seg = str(x.get("no_segmento", "")).lower()
        if seg.startswith("disp") or seg.startswith("im"):
            continue
        total += float(x.get("vl_total_atual") or 0)
    return round(total, 2) if linhas else None


def montar(planilhas, dados, ano_atual, ano_anterior, antigos):
    atual = dados[ano_atual]
    ranking, total_rank = calcular_ranking(atual)
    cnpjs = sorted({v["cnpj"] for v in atual.values() if v["cnpj"]})

    crps, ultimos = {}, {}
    print("CRP: buscando...", flush=True)
    try:
        crps = buscar_crp(cnpjs)
    except Exception as e:  # noqa: BLE001
        print("::warning::CRP: API do CADPREV indisponível, mantendo valores anteriores:", e)
    print(f"CRP ok: {len(crps)}. DAIR: buscando último mês enviado...", flush=True)
    try:
        ultimos = ultimo_mes_dair(date.today().year)
    except Exception as e:  # noqa: BLE001
        print("::warning::DAIR: API do CADPREV indisponível, mantendo valores anteriores:", e)

    # saldo: só rebusca quem tem mês novo desde a última execução
    a_buscar = []
    for k, v in atual.items():
        c = v["cnpj"]
        ant = antigos.get(k, {})
        if c in ultimos:
            ano, mes = ultimos[c]
            ref = f"{ano}-{mes:02d}"
            if ant.get("sm") == ref and ant.get("s") is not None:
                v["s"], v["sm"] = ant["s"], ref
            else:
                a_buscar.append((k, c, ano, mes, ref))
        elif ant.get("s") is not None:
            v["s"], v["sm"] = ant["s"], ant.get("sm", "")

    def um(item):
        k, c, ano, mes, ref = item
        try:
            return k, saldo_dair(c, ano, mes), ref
        except Exception:  # noqa: BLE001
            return k, None, ref
    print(f"Saldo DAIR: {len(a_buscar)} entes para buscar.")
    feitos = 0
    with ThreadPoolExecutor(12) as ex:
        for k, s, ref in ex.map(um, a_buscar):
            feitos += 1
            if feitos % 200 == 0:
                print(f"  saldo: {feitos}/{len(a_buscar)}", flush=True)
            if s is not None:
                atual[k]["s"], atual[k]["sm"] = s, ref
            elif antigos.get(k, {}).get("s") is not None:
                atual[k]["s"], atual[k]["sm"] = antigos[k]["s"], antigos[k].get("sm", "")

    entes = []
    for (uf, nome), d in sorted(atual.items()):
        ant = dados[ano_anterior].get((uf, nome), {})
        e = {
            "n": nome, "uf": uf,
            "g": d["grupo"], "sg": d["subgrupo"], "isp": d["isp"], "pa": d["perfil"],
            "gAnt": ant.get("grupo", ""), "pg": d.get("pg", "-"),
        }
        if (uf, nome) in ranking:
            e["rk"] = ranking[(uf, nome)]
        if "s" in d:
            e["s"], e["sm"] = d["s"], d["sm"]
        crp = crps.get(d["cnpj"]) or antigos.get((uf, nome), {}).get("crp")
        if crp:
            e["crp"] = crp
        entes.append(e)
    return {
        "meta": {
            "fonte": "Ministério da Previdência Social (SPREV): ISP-RPPS; CADPREV: DAIR e CRP",
            "paginaOficial": PAGINA_ISP,
            "anoISP": ano_atual, "anoISPAnterior": ano_anterior,
            "urlPlanilha": planilhas[ano_atual],
            "totalRanking": total_rank,
            "geradoEm": date.today().isoformat(),
        },
        "entes": entes,
    }


def principal():
    html_alvo = Path(sys.argv[sys.argv.index("--html") + 1]) if "--html" in sys.argv else None
    planilhas = achar_planilhas_finais()
    anos = sorted(planilhas)
    if len(anos) < 2:
        sys.exit(f"Não achei duas planilhas finais na página do ISP: {planilhas}")
    ano_atual, ano_anterior = anos[-1], anos[-2]
    tmp = Path(tempfile.mkdtemp(prefix="isp_"))
    dados = {}
    for ano in (ano_atual, ano_anterior):
        arq = tmp / f"isp{ano}.xlsx"
        print("Baixando", ano, planilhas[ano])
        baixar(planilhas[ano], arq)
        dados[ano] = ler_planilha(arq)

    # dados da execução anterior (para reaproveitar saldo e não perder nada se uma API falhar)
    if html_alvo:
        achou = BLOCO.search(html_alvo.read_text(encoding="utf-8"))
        antigo = json.loads(achou.group(2).replace(FECHA, "</")) if achou else {}
    else:
        antigo = json.loads(DESTINO.read_text(encoding="utf-8")) if DESTINO.exists() else {}
    antigos = {(e["uf"], e["n"]): e for e in antigo.get("entes", [])}

    novo = montar(planilhas, dados, ano_atual, ano_anterior, antigos)

    meta_antiga = antigo.get("meta", {})
    if antigo.get("entes") == novo["entes"] and meta_antiga.get("urlPlanilha") == planilhas[ano_atual]:
        print("Sem mudanças nos dados.")
    else:
        texto = json.dumps(novo, ensure_ascii=False, separators=(",", ":"))
        if html_alvo:
            html = html_alvo.read_text(encoding="utf-8")
            seguro = texto.replace("</", FECHA)
            html = BLOCO.sub(lambda m: m.group(1) + seguro + m.group(3), html, count=1)
            html_alvo.write_text(html, encoding="utf-8")
            print(f"OK: {len(novo['entes'])} entes gravados em {html_alvo}")
        else:
            DESTINO.parent.mkdir(parents=True, exist_ok=True)
            DESTINO.write_text(texto, encoding="utf-8")
            print(f"OK: {len(novo['entes'])} entes gravados em {DESTINO}")

    v = versao_do_manual()
    if v and v > MANUAL_NO_APP:
        msg = f"Saiu o Manual de Certificação versão {v[0]}.{v[1]}. Revise regras.js (hoje usa {MANUAL_NO_APP[0]}.{MANUAL_NO_APP[1]})."
        print("::warning::" + msg)
        Path("ALERTA-NORMA.txt").write_text(msg + "\n", encoding="utf-8")
        sys.exit(3)  # faz o GitHub avisar por e-mail que o job falhou


if __name__ == "__main__":
    principal()
